"""Read-only OOXML evidence reader. Formula caches are observations, never evaluated truth."""
import hashlib
import json
import re
import warnings
from collections import Counter
from datetime import date, datetime
from pathlib import Path
import openpyxl
from .normalizers import text


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def stable_id(*parts):
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:24]


def scalar(v):
    return v.isoformat() if isinstance(v, (date, datetime)) else v


def inspect(path):
    path = Path(path)
    checksum = digest(path)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        wb = openpyxl.load_workbook(path, data_only=False)
        cached = openpyxl.load_workbook(path, data_only=True)
    sheets, rows = [], []
    for ws in wb:
        formulas, headers, nonempty, useful = [], [], [], 0
        types = Counter()
        for cells in ws:
            populated = [c for c in cells if c.value is not None]
            if not populated:
                continue
            useful += any(c.data_type != 'f' for c in populated)
            if cells[0].row <= 10:
                labels = {c.coordinate: scalar(c.value) for c in populated if c.data_type != 'f'}
                if len(labels) >= 2:
                    headers.append({'row': cells[0].row, 'labels': labels})
            refs = []
            for c in populated:
                types[c.data_type] += 1
                raw = scalar(c.value)
                is_formula = c.data_type == 'f'
                ref = {'source_file': path.name, 'workbook_sha256': checksum,
                       'source_sheet': ws.title, 'source_row': c.row, 'source_column': c.column,
                       'source_cell': c.coordinate, 'raw_value': raw,
                       'normalized_value': None if is_formula else text(raw),
                       'data_type': c.data_type, 'formula': raw if is_formula else None,
                       'cached_value': scalar(cached[ws.title][c.coordinate].value) if is_formula else None,
                       'number_format': c.number_format}
                refs.append(ref)
                if is_formula:
                    formulas.append({'cell': c.coordinate, 'formula': raw, 'cached_value': ref['cached_value'],
                                     'sheet_references': sorted(set(re.findall(r"(?:'([^']+)'|([A-Za-z_][\w]*))!", raw)))})
            rows.append({'source_sheet': ws.title, 'source_row': cells[0].row, 'cells': refs})
            nonempty.append(cells[0].row)
        sheets.append({'name': ws.title, 'state': ws.sheet_state, 'dimensions': ws.calculate_dimension(),
                       'max_row': ws.max_row, 'max_column': ws.max_column,
                       'nonempty_rows': len(nonempty), 'literal_rows': useful,
                       'header_candidates': headers, 'data_types': dict(types),
                       'formula_count': len(formulas), 'formulas': formulas,
                       'merged_cells': [str(r) for r in ws.merged_cells.ranges],
                       'hidden_rows': [i for i, d in ws.row_dimensions.items() if d.hidden],
                       'hidden_columns': [i for i, d in ws.column_dimensions.items() if d.hidden],
                       'validations': [{'range': str(v.sqref), 'type': v.type,
                                        'formula1': v.formula1, 'formula2': v.formula2}
                                       for v in ws.data_validations.dataValidation]})
    result = {'source_file': path.name, 'sha256': checksum,
              'workbook_properties': {'title': wb.properties.title, 'version': wb.properties.version},
              'reader_warnings': sorted(set(str(w.message) for w in caught)),
              'defined_names': {k: str(v.attr_text) for k, v in wb.defined_names.items()}, 'sheets': sheets}
    wb.close()
    cached.close()
    if digest(path) != checksum:
        raise RuntimeError('Source changed during read: ' + path.name)
    return result, rows
